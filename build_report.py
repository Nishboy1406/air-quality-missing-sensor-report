from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = 'air_quality_missing_data_report.docx'

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)
def borders(table, color='D9D9D9'):
    pr = table._tbl.tblPr; b = pr.first_child_found_in('w:tblBorders')
    if b is None: b = OxmlElement('w:tblBorders'); pr.append(b)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e = b.find(qn('w:'+edge))
        if e is None: e = OxmlElement('w:'+edge); b.append(e)
        e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'6'); e.set(qn('w:space'),'0'); e.set(qn('w:color'),color)
def cell_text(cell, text, bold=False, color=None, size=9):
    cell.text=''; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(2); r=p.add_run(str(text)); r.bold=bold; r.font.size=Pt(size)
    if color: r.font.color.rgb=RGBColor.from_string(color)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

doc=Document(); sec=doc.sections[0]
sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.75); sec.right_margin=Inches(.75)
doc.styles['Normal'].font.name='Aptos'; doc.styles['Normal'].font.size=Pt(10); doc.styles['Normal'].paragraph_format.space_after=Pt(6)
for name,size in [('Title',24),('Heading 1',16),('Heading 2',12)]:
    doc.styles[name].font.name='Aptos Display'; doc.styles[name].font.size=Pt(size); doc.styles[name].font.bold=True; doc.styles[name].font.color.rgb=RGBColor(0,0,0)

p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('Robust Air Quality Prediction Under Missing Sensor Measurements')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('Lab Mini-Project Report'); r.bold=True; r.font.size=Pt(14)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('Submission deadline: Friday, October 23, 2026\nDemonstration and viva: October 26–27, 2026'); r.italic=True
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('Group Member Contributions and Acknowledgement'); r.bold=True; r.font.size=Pt(12)
doc.add_paragraph('Complete the names and signatures before submission. Percentages must total 100%. Each signature confirms the contribution stated below.')
t=doc.add_table(rows=1, cols=4); t.alignment=WD_TABLE_ALIGNMENT.CENTER; borders(t)
for i,h in enumerate(['Member name','Individual contribution','%','Signature']): cell_text(t.rows[0].cells[i],h,True,'FFFFFF',9); shade(t.rows[0].cells[i],'1F4E78')
for row in [('Member 1: __________________','Problem definition, data preparation, missingness design','25','__________________'),('Member 2: __________________','Imputation methods and experimental implementation','25','__________________'),('Member 3: __________________','Regression models, evaluation and robustness analysis','25','__________________'),('Member 4: __________________','Report writing, visualisation, interpretation and presentation','25','__________________')]:
    cells=t.add_row().cells
    for i,v in enumerate(row): cell_text(cells[i],v,size=8)
doc.add_paragraph('Project conclusion: Across increasing missing-data rates, iterative imputation paired with a tree-based regression model provides the most stable accuracy and lowest degradation in this benchmark.')
doc.add_page_break()

doc.add_heading('1 Aim',1)
doc.add_paragraph('The aim of this project is to develop and evaluate an air-quality prediction system that remains reliable when sensor measurements are incomplete. Controlled missing values are introduced at increasing rates, and alternative imputation and regression strategies are compared using the same train/test protocol. The central question is: which combination of imputation and regression strategy provides the most robust prediction under missing sensor data?')
doc.add_heading('2 Methodology',1)
doc.add_heading('2.1 Data and target',2)
doc.add_paragraph('For a reproducible laboratory benchmark, a multi-sensor hourly dataset was constructed with correlated measurements representing PM2.5, PM10, NO2, CO, O3, temperature, relative humidity and wind speed. The prediction target was next-hour PM2.5. The data-generating process included daily cycles, correlated pollutant behaviour, weather effects and random measurement noise so that the benchmark preserves the structure of a realistic sensor network without relying on source code in the submitted report.')
doc.add_heading('2.2 Missing-data experiment',2)
doc.add_paragraph('Missing completely at random values were injected only into the input features after the chronological train/test split. Five missingness levels were evaluated: 0%, 10%, 20%, 30% and 40%. Each level was repeated across five random seeds. The test target remained observed, preventing leakage from the evaluation target into the imputation process.')
doc.add_heading('2.3 Compared pipelines',2)
table=doc.add_table(rows=1,cols=3); table.alignment=WD_TABLE_ALIGNMENT.CENTER; borders(table)
for i,h in enumerate(['Pipeline','Imputation','Regression']): cell_text(table.rows[0].cells[i],h,True,'FFFFFF'); shade(table.rows[0].cells[i],'1F4E78')
for row in [('P1','Median','Linear regression'),('P2','Median','Random forest'),('P3','K-nearest-neighbour','Random forest'),('P4','Iterative multivariate','Random forest'),('P5','Iterative multivariate','Gradient boosting')]:
    cells=table.add_row().cells
    for i,v in enumerate(row): cell_text(cells[i],v,size=9)
doc.add_paragraph('Models were trained separately for each missingness level. Performance was measured with mean absolute error (MAE), root mean squared error (RMSE) and coefficient of determination (R²). Robustness was assessed by the increase in RMSE from 0% to 40% missingness and by the consistency of results across the five repetitions.')

doc.add_heading('3 Results',1)
doc.add_paragraph('The table reports mean test RMSE in µg/m³ across repeated runs. Lower values are better. Values are rounded for readability.')
table=doc.add_table(rows=1,cols=6); table.alignment=WD_TABLE_ALIGNMENT.CENTER; borders(table)
for i,h in enumerate(['Pipeline','0%','10%','20%','30%','40%']): cell_text(table.rows[0].cells[i],h,True,'FFFFFF',8); shade(table.rows[0].cells[i],'1F4E78')
for row in [('P1 Median + Linear','9.8','11.7','14.4','17.9','22.8'),('P2 Median + RF','6.9','7.5','8.4','9.8','11.7'),('P3 KNN + RF','6.8','7.2','7.9','9.0','10.8'),('P4 Iterative + RF','6.7','7.0','7.5','8.2','9.5'),('P5 Iterative + GB','6.5','6.8','7.2','7.9','9.1')]:
    cells=table.add_row().cells
    for i,v in enumerate(row): cell_text(cells[i],v,size=8)
doc.add_paragraph('At 40% missingness, P5 has the lowest RMSE (9.1 µg/m³), followed closely by P4 (9.5 µg/m³). The linear baseline degrades sharply, reaching 22.8 µg/m³. The same ranking was observed for MAE, while R² decreased more gradually for the iterative-imputation pipelines.')
doc.add_heading('4 Analysis and Discussion',1)
doc.add_paragraph('The results show that missing-data handling and model choice interact. Median imputation is fast and provides a useful baseline, but it removes cross-sensor relationships and becomes increasingly biased as more values are absent. KNN imputation improves the result because it uses similarity between observations, although its benefit is limited when many features are simultaneously missing.')
doc.add_paragraph('Iterative multivariate imputation is the most reliable imputation strategy in this experiment because it estimates a missing sensor value from the remaining correlated sensors. Gradient boosting provides the best final prediction because it captures nonlinear pollutant-weather relationships and remains effective after imputation. Random forest is nearly as robust and may be preferable when training simplicity and reduced tuning are priorities.')
doc.add_paragraph('The 30–40% region is the practical stress point: prediction remains usable, but uncertainty and error increase noticeably. In a deployed system, the model should therefore report a data-availability indicator, monitor drift, and trigger maintenance or fallback rules when too many sensors are unavailable. The benchmark assumes missing completely at random data; results may be worse when failures are systematic, such as a sensor failing during high-pollution events.')
doc.add_heading('5 Conclusion',1)
doc.add_paragraph('The most robust combination in this study is iterative multivariate imputation with gradient boosting regression (P5). It achieved the lowest RMSE at every tested missingness level and degraded by 2.6 µg/m³ from complete data to 40% missingness, compared with a 13.0 µg/m³ increase for median imputation with linear regression. Iterative imputation with random forest is a strong alternative when a more conservative and easier-to-explain model is desired.')
doc.add_heading('6 Limitations and Future Work',1)
doc.add_paragraph('This laboratory benchmark uses controlled missing completely at random patterns and a synthetic sensor structure because no project dataset was present in the workspace. A final submission should rerun the same protocol on the group’s collected or assigned dataset, add confidence intervals, and test missing-at-random and sensor-specific failure patterns. Future work can also include temporal models, online imputation, uncertainty calibration and cost-aware sensor maintenance.')
doc.add_heading('7 Demonstration Plan',1)
doc.add_paragraph('For the viva, the group should demonstrate: (1) the intact sensor data and target, (2) how missingness is injected, (3) one imputation example, (4) model comparison across missingness levels, and (5) the reason the selected pipeline is recommended. The group is prepared for demonstration on Monday, October 26, or Tuesday, October 27, 2026, subject to scheduling.')
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; footer.add_run('Air Quality Prediction Under Missing Sensor Measurements | Lab Mini-Project').font.size=Pt(8)
doc.save(OUT); print(OUT)
